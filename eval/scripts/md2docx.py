"""Minimal Markdown -> .docx converter (headings, paragraphs, lists, tables)."""

from __future__ import annotations

import re
import sys
from pathlib import Path

import docx
from docx.shared import Pt


def add_runs(par, text: str):
    parts = re.split(r"(\*\*.+?\*\*|`[^`]+`)", text)
    for p in parts:
        if p.startswith("**") and p.endswith("**"):
            r = par.add_run(p[2:-2]); r.bold = True
        elif p.startswith("`") and p.endswith("`"):
            r = par.add_run(p[1:-1])
            r.font.name = "Consolas"; r.font.size = Pt(9)
        elif p:
            par.add_run(p)


def convert(md_path: Path, out_path: Path) -> None:
    doc = docx.Document()
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    table_buf = []
    while i < len(lines):
        line = lines[i]
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s\-:|]+\|$", lines[i + 1]):
            while i < len(lines) and line.startswith("|"):
                table_buf.append([c.strip() for c in line.strip("|").split("|")])
                i += 1
                line = lines[i] if i < len(lines) else ""
            if table_buf:
                rows = len(table_buf)
                cols = max(len(r) for r in table_buf)
                t = doc.add_table(rows=rows, cols=cols)
                t.style = "Light Grid Accent 1"
                for r, row in enumerate(table_buf):
                    for c in range(cols):
                        text = row[c] if c < len(row) else ""
                        add_runs(t.cell(r, c).paragraphs[0], text)
                table_buf = []
            continue
        if line.startswith("|"):
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif re.match(r"^!\[.*?\]\(.*?\)\s*$", line):
            m = re.match(r"^!\[(.*?)\]\((.*?)\)\s*$", line)
            path = Path(m.group(2))
            if not path.is_absolute():
                path = (md_path.parent / path).resolve()
            if path.exists():
                try:
                    doc.add_picture(str(path), width=docx.shared.Inches(6.3))
                except Exception as exc:  # noqa: BLE001
                    doc.add_paragraph(f"[image failed: {path} ({exc})]")
            else:
                doc.add_paragraph(f"[missing image: {path}]")
            if m.group(1):
                cap = doc.add_paragraph()
                r = cap.add_run(m.group(1)); r.italic = True; r.font.size = Pt(9)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:])
        elif re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\.\s", "", line))
        elif line.startswith("---"):
            doc.add_paragraph()
        else:
            p = doc.add_paragraph()
            add_runs(p, line)
        i += 1
    doc.save(str(out_path))
    print(f"wrote {out_path}")


if __name__ == "__main__":
    src = Path(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".docx")
    convert(src, out)